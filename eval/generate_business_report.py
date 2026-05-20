"""
eval/generate_business_report.py

Generates a business-friendly Excel report comparing two eval runs.
Sheets: Summary (first) | CUSTOM | CANVAS

Each detail sheet columns (when --after is provided):
  Project Name | Original Narrative | New Score (/10) | AI Rewritten Narrative | AI Reasoning

Usage:
    cd eval
    python generate_business_report.py --before baseline --after improved
    python generate_business_report.py --before baseline --after improved --system rag
    python generate_business_report.py --before baseline --after improved --system agent

Environment:
    Reads result files produced by run_eval.py:
        results_<run>_rag.json
        results_<run>_agent.json
    Reads ground_truth.json produced by extract_ground_truth.py.
"""

import argparse
import json
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE    = pathlib.Path(__file__).parent
GT_PATH = HERE / "ground_truth.json"

# ── Colours ───────────────────────────────────────────────────────────────────
C_NAVY        = "1F3864"
C_TITLE_BG    = "2E75B6"
C_GREEN       = "C6EFCE"
C_YELLOW      = "FFEB9C"
C_RED         = "FFC7CE"
C_LTBLUE      = "DDEBF7"
C_SCORE_HIGH  = "C6EFCE"
C_SCORE_MID   = "FFEB9C"
C_SCORE_LOW   = "FFC7CE"
C_ROW_ALT     = "F2F7FB"
C_WHITE       = "FFFFFF"
C_BORDER      = "9DC3E6"

# ── System name mapping ───────────────────────────────────────────────────────
_SHEET_NAMES   = {"rag": "CUSTOM",          "agent": "CANVAS"}
_DISPLAY_NAMES = {"rag": "Custom Approach", "agent": "Canvas App Approach"}


# ── Style helpers ─────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, white=False, size=11, italic=False) -> Font:
    return Font(bold=bold, italic=italic, size=size,
                color="FFFFFF" if white else "000000")


def _border() -> Border:
    s = Side(border_style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


def _cell(ws, row: int, col: int, value,
          bg=C_WHITE, bold=False, white=False,
          center=False, wrap=True, size=11, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.fill      = _fill(bg)
    c.font      = _font(bold=bold, white=white, size=size, italic=italic)
    c.border    = _border()
    c.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=wrap,
    )
    return c


def _score_bg(score) -> str:
    if score is None: return C_WHITE
    if score >= 8:    return C_SCORE_HIGH
    if score >= 6:    return C_SCORE_MID
    return C_SCORE_LOW


def _verdict(score) -> str:
    if score is None: return "—"
    if score >= 8:    return "PASS"
    if score >= 6:    return "WARNING"
    return "FAIL"


def _delta_str(a, b) -> str:
    if a is None or b is None: return "—"
    d = round(b - a, 1)
    return f"+{d}" if d > 0 else str(d)


def _delta_bg(a, b) -> str:
    if a is None or b is None: return C_WHITE
    d = b - a
    if d > 0: return C_GREEN
    if d < 0: return C_RED
    return C_LTBLUE


# ── Data loaders ──────────────────────────────────────────────────────────────

def _load_results(run_name: str, system: str) -> dict:
    path = HERE / f"results_{run_name}_{system}.json"
    if not path.exists():
        return {}
    entries = json.loads(path.read_text(encoding="utf-8"))
    return {e["project_name"]: e for e in entries}


def _score_from_entry(entry: dict):
    if not entry or entry.get("status") != "ok":
        return None
    return entry.get("result", {}).get("layer1", {}).get("compliance_score")


def _rewritten_from_entry(entry: dict) -> str:
    if not entry or entry.get("status") != "ok":
        return "—"
    return entry.get("result", {}).get("rewritten_narrative") or "—"


def _reasoning_from_entry(entry: dict) -> str:
    if not entry or entry.get("status") != "ok":
        return "—"
    issues = entry.get("result", {}).get("layer1", {}).get("issues", [])
    if not issues:
        return "All criteria passed."
    return "\n".join(f"• {issue}" for issue in issues)


# ── Summary stats ─────────────────────────────────────────────────────────────

def _stats(rows: list) -> dict:
    before = [r["score_before"] for r in rows if r["score_before"] is not None]
    after  = [r["score_after"]  for r in rows if r["score_after"]  is not None]
    improved  = sum(1 for r in rows if r["score_before"] is not None and r["score_after"] is not None and r["score_after"] > r["score_before"])
    declined  = sum(1 for r in rows if r["score_before"] is not None and r["score_after"] is not None and r["score_after"] < r["score_before"])
    no_change = sum(1 for r in rows if r["score_before"] is not None and r["score_after"] is not None and r["score_after"] == r["score_before"])
    return {
        "avg_before": round(sum(before) / len(before), 1) if before else None,
        "avg_after":  round(sum(after)  / len(after),  1) if after  else None,
        "improved": improved, "declined": declined, "no_change": no_change,
        "pass_b": sum(1 for r in rows if _verdict(r["score_before"]) == "PASS"),
        "warn_b": sum(1 for r in rows if _verdict(r["score_before"]) == "WARNING"),
        "fail_b": sum(1 for r in rows if _verdict(r["score_before"]) == "FAIL"),
        "pass_a": sum(1 for r in rows if _verdict(r["score_after"])  == "PASS"),
        "warn_a": sum(1 for r in rows if _verdict(r["score_after"])  == "WARNING"),
        "fail_a": sum(1 for r in rows if _verdict(r["score_after"])  == "FAIL"),
    }


# ── Summary sheet ─────────────────────────────────────────────────────────────

def _build_summary_sheet(wb, custom_rows: list, canvas_rows: list,
                         before_name: str, after_name: str) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False

    row = 1

    t = ws.cell(row=row, column=1, value="NDA Narrative Validation — Evaluation Summary")
    t.font      = Font(bold=True, size=16, color=C_NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"A{row}:G{row}")
    ws.row_dimensions[row].height = 38
    row += 1

    n = max(len(custom_rows), len(canvas_rows))
    sub = ws.cell(row=row, column=1,
        value=(f"Comparing {before_name} (old guidance) vs {after_name} (new guidance)"
               f" — {n} projects evaluated per system"))
    sub.font      = Font(size=11, italic=True, color="595959")
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"A{row}:G{row}")
    ws.row_dimensions[row].height = 22
    row += 2

    rs  = _stats(custom_rows)
    as_ = _stats(canvas_rows)

    def _section(title):
        nonlocal row
        sec = ws.cell(row=row, column=1, value=title)
        sec.font      = Font(bold=True, size=12, color="FFFFFF")
        sec.fill      = _fill(C_NAVY)
        sec.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(f"A{row}:G{row}")
        ws.row_dimensions[row].height = 26
        row += 1

    # ── Section 1: Average Scores ─────────────────────────────────────────────
    _section("Average Compliance Score (out of 10)")
    for ci, h in enumerate(["System", "Old Guidance", "New Guidance", "Change"], 1):
        _cell(ws, row, ci, h, bg=C_TITLE_BG, bold=True, white=True, center=True)
    ws.row_dimensions[row].height = 22
    row += 1

    for label, s in [("Custom Approach", rs), ("Canvas App Approach", as_)]:
        _cell(ws, row, 1, label, bold=True)
        _cell(ws, row, 2, s["avg_before"] if s["avg_before"] is not None else "—", center=True)
        _cell(ws, row, 3, s["avg_after"]  if s["avg_after"]  is not None else "—", center=True)
        _cell(ws, row, 4, _delta_str(s["avg_before"], s["avg_after"]),
              bg=_delta_bg(s["avg_before"], s["avg_after"]), center=True, bold=True)
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1

    # ── Section 2: Score Movement ─────────────────────────────────────────────
    _section("Score Movement Across Projects")
    for ci, h in enumerate(["System", "Improved", "No Change", "Declined"], 1):
        _cell(ws, row, ci, h, bg=C_TITLE_BG, bold=True, white=True, center=True)
    ws.row_dimensions[row].height = 22
    row += 1

    for label, s in [("Custom Approach", rs), ("Canvas App Approach", as_)]:
        _cell(ws, row, 1, label, bold=True)
        _cell(ws, row, 2, s["improved"],  bg=C_GREEN  if s["improved"]  > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 3, s["no_change"], bg=C_LTBLUE if s["no_change"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 4, s["declined"],  bg=C_RED    if s["declined"]  > 0 else C_WHITE, center=True, bold=True)
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1

    # ── Section 3: Verdict Breakdown ──────────────────────────────────────────
    _section("Verdict Breakdown — Old Guidance vs New Guidance")
    for ci, h in enumerate(
        ["System", "PASS (Old)", "WARNING (Old)", "FAIL (Old)",
         "PASS (New)", "WARNING (New)", "FAIL (New)"], 1):
        _cell(ws, row, ci, h, bg=C_TITLE_BG, bold=True, white=True, center=True)
    ws.row_dimensions[row].height = 22
    row += 1

    for label, s in [("Custom Approach", rs), ("Canvas App Approach", as_)]:
        _cell(ws, row, 1, label, bold=True)
        _cell(ws, row, 2, s["pass_b"], bg=C_GREEN  if s["pass_b"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 3, s["warn_b"], bg=C_YELLOW if s["warn_b"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 4, s["fail_b"], bg=C_RED    if s["fail_b"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 5, s["pass_a"], bg=C_GREEN  if s["pass_a"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 6, s["warn_a"], bg=C_YELLOW if s["warn_a"] > 0 else C_WHITE, center=True, bold=True)
        _cell(ws, row, 7, s["fail_a"], bg=C_RED    if s["fail_a"] > 0 else C_WHITE, center=True, bold=True)
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1

    # ── Section 4: Per-Project Quick View ─────────────────────────────────────
    _section("Per-Project Quick View")
    for ci, h in enumerate(
        ["Project", "Custom Old", "Custom New", "Custom Change",
         "Canvas Old", "Canvas New", "Canvas Change"], 1):
        _cell(ws, row, ci, h, bg=C_TITLE_BG, bold=True, white=True, center=True)
    ws.row_dimensions[row].height = 22
    row += 1

    for i, (rp, ap) in enumerate(zip(custom_rows, canvas_rows)):
        bg = "F2F7FB" if i % 2 == 0 else C_WHITE
        rb, ra = rp["score_before"], rp["score_after"]
        ab, aa = ap["score_before"], ap["score_after"]
        _cell(ws, row, 1, rp["name"], bg=bg, bold=True)
        _cell(ws, row, 2, rb if rb is not None else "—", bg=_score_bg(rb), center=True, bold=True)
        _cell(ws, row, 3, ra if ra is not None else "—", bg=_score_bg(ra), center=True, bold=True)
        _cell(ws, row, 4, _delta_str(rb, ra), bg=_delta_bg(rb, ra), center=True, bold=True)
        _cell(ws, row, 5, ab if ab is not None else "—", bg=_score_bg(ab), center=True, bold=True)
        _cell(ws, row, 6, aa if aa is not None else "—", bg=_score_bg(aa), center=True, bold=True)
        _cell(ws, row, 7, _delta_str(ab, aa), bg=_delta_bg(ab, aa), center=True, bold=True)
        ws.row_dimensions[row].height = 22
        row += 1

    row += 1
    _cell(ws, row, 1, "Score legend:", bold=True)
    _cell(ws, row, 2, "8-10 = PASS",   bg=C_GREEN,  center=True)
    _cell(ws, row, 3, "6-7 = WARNING", bg=C_YELLOW, center=True)
    _cell(ws, row, 4, "0-5 = FAIL",    bg=C_RED,    center=True)

    ws.column_dimensions["A"].width = 44
    for col_letter in ["B", "C", "D", "E", "F", "G"]:
        ws.column_dimensions[col_letter].width = 15


# ── Detail sheet ──────────────────────────────────────────────────────────────

def _build_sheet(wb, system: str, ground_truth: list,
                 before_name: str, after_name: str | None) -> list:
    """
    Build the CUSTOM or CANVAS detail sheet.
    Returns a list of {name, score_before, score_after} dicts for the summary.
    """
    sheet_name   = _SHEET_NAMES.get(system,   system.upper())
    display_name = _DISPLAY_NAMES.get(system, system.upper())

    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    before_map = _load_results(before_name, system)
    after_map  = _load_results(after_name,  system) if after_name else {}
    has_after  = bool(after_map)

    # ── Title ─────────────────────────────────────────────────────────────────
    n_cols = 5 if has_after else 3
    title  = (f"Narrative Validation Results — {display_name} — "
               f"{before_name.title()} vs {after_name.title()}"
              if has_after else
               f"Narrative Validation Results — {display_name} — {before_name.title()}")

    row = 1
    tc = ws.cell(row=row, column=1, value=title)
    tc.font      = Font(bold=True, size=13, color="FFFFFF")
    tc.fill      = _fill(C_TITLE_BG)
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    ws.row_dimensions[row].height = 28
    row += 1

    # ── Column headers ────────────────────────────────────────────────────────
    if has_after:
        headers = [
            "Project Name",
            "Original Narrative",
            "New Score (/10)",
            "AI Rewritten Narrative",
            "AI Reasoning (Why It Failed / Passed)",
        ]
    else:
        headers = ["Project Name", "Narrative", "Score (/10)"]

    for ci, h in enumerate(headers, 1):
        _cell(ws, row, ci, h, bg=C_NAVY, bold=True, white=True, center=True)
    ws.row_dimensions[row].height = 36
    row += 1

    # ── Data rows ─────────────────────────────────────────────────────────────
    score_rows = []
    for i, gt in enumerate(ground_truth):
        name      = gt["project_name"]
        narrative = gt.get("original_narrative", "")
        row_bg    = C_ROW_ALT if i % 2 == 0 else C_WHITE

        before_entry = before_map.get(name, {})
        after_entry  = after_map.get(name,  {})

        score_before = _score_from_entry(before_entry)
        score_after  = _score_from_entry(after_entry) if has_after else None
        score_rows.append({"name": name, "score_before": score_before, "score_after": score_after})

        narrative_display = narrative[:1000] + ("…" if len(narrative) > 1000 else "")
        _cell(ws, row, 1, name, bg=row_bg, bold=True)
        _cell(ws, row, 2, narrative_display, bg=row_bg)

        if has_after:
            _cell(ws, row, 3,
                  score_after if score_after is not None else "—",
                  bg=_score_bg(score_after), center=True, bold=True)

            rewritten = _rewritten_from_entry(after_entry)
            rewritten_display = rewritten[:2000] + ("…" if len(rewritten) > 2000 else "")
            _cell(ws, row, 4, rewritten_display, bg=row_bg)
            _cell(ws, row, 5, _reasoning_from_entry(after_entry), bg=row_bg)
        else:
            _cell(ws, row, 3,
                  score_before if score_before is not None else "—",
                  bg=_score_bg(score_before), center=True, bold=True)

        ws.row_dimensions[row].height = 120
        row += 1

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 16
    if has_after:
        ws.column_dimensions["D"].width = 65
        ws.column_dimensions["E"].width = 55

    # ── Legend ────────────────────────────────────────────────────────────────
    row += 1
    _cell(ws, row, 1, "Score legend:", bold=True)
    _cell(ws, row, 2, "8–10 = PASS",   bg=C_SCORE_HIGH, center=True)
    _cell(ws, row, 3, "6–7 = WARNING", bg=C_SCORE_MID,  center=True)
    if has_after:
        _cell(ws, row, 4, "0–5 = FAIL", bg=C_SCORE_LOW, center=True)

    return score_rows


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Generate a business-friendly before/after narrative eval report."
    )
    parser.add_argument("--before",  default="baseline",
                        help="Run name for 'before' results (default: baseline)")
    parser.add_argument("--after",   default=None,
                        help="Run name for 'after' results (omit for single-run report)")
    parser.add_argument("--system",  choices=["rag", "agent", "both"], default="both",
                        help="Which system(s) to include (default: both)")
    args = parser.parse_args()

    if not GT_PATH.exists():
        raise SystemExit(f"ERROR: {GT_PATH} not found.\nRun extract_ground_truth.py first.")

    ground_truth: list = json.loads(GT_PATH.read_text(encoding="utf-8"))
    systems = ["rag", "agent"] if args.system == "both" else [args.system]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    score_data = {}
    for system in systems:
        rows = _build_sheet(wb, system, ground_truth, args.before, args.after)
        score_data[system] = rows

    if args.after and "rag" in score_data and "agent" in score_data:
        _build_summary_sheet(wb, score_data["rag"], score_data["agent"],
                             args.before, args.after)

    suffix = f"{args.before}_vs_{args.after}" if args.after else args.before
    out    = HERE / f"business_report_{suffix}.xlsx"
    wb.save(out)
    print(f"Saved -> {out}")


if __name__ == "__main__":
    main()