"""
eval/generate_excel_report.py  —  Phase 5

Generates a before/after Excel workbook comparing two eval runs.

Sheets:
  1. Summary        — overall metrics side by side for both systems
  2. Per-Project    — score, verdict, delta per project for each system
  3. Per-Criterion  — recall/precision/F1 before and after per criterion

Usage:
    cd eval
    python generate_excel_report.py                          # baseline only
    python generate_excel_report.py --before baseline --after improved
    python generate_excel_report.py --before baseline --after improved --system rag
"""

import argparse
import json
import pathlib

import openpyxl
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side
)
from openpyxl.utils import get_column_letter

HERE    = pathlib.Path(__file__).parent
GT_PATH = HERE / "ground_truth.json"

# ── Colours ───────────────────────────────────────────────────────────────────
C_HEADER    = "1F3864"   # dark navy
C_SUBHEADER = "2E75B6"   # medium blue
C_PASS      = "C6EFCE"   # green fill
C_WARN      = "FFEB9C"   # yellow fill
C_FAIL      = "FFC7CE"   # red fill
C_NEUTRAL   = "DDEBF7"   # light blue fill
C_WHITE     = "FFFFFF"
C_BORDER    = "9DC3E6"

CRITERIA_ORDER = [
    "project_description",
    "project_benefit",
    "dca_rag",
    "completion_cost",
    "schedule",
    "baseline_rag",
    "baseline_movement",
    "highlights_issues",
    "cap_cap_rag",
]

CRITERIA_LABELS = {
    "project_description": "Project Description",
    "project_benefit":     "Project Benefit",
    "dca_rag":             "DCA RAG",
    "completion_cost":     "Completion Cost (P50/P80)",
    "schedule":            "Schedule Position",
    "baseline_rag":        "Baseline RAG",
    "baseline_movement":   "Baseline Movement",
    "highlights_issues":   "Highlights / Issues",
    "cap_cap_rag":         "Cap/Cap RAG",
}

VERDICT_FILL = {
    "PASS": C_PASS,
    "PASS_WITH_WARNINGS": C_WARN,
    "WARN": C_WARN,
    "FAIL": C_FAIL,
}


# ── Style helpers ─────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, white=False, size=11) -> Font:
    color = "FFFFFF" if white else "000000"
    return Font(bold=bold, color=color, size=size)


def _border() -> Border:
    s = Side(border_style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


def _header_cell(ws, row, col, value, span=1, bg=C_HEADER):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill    = _fill(bg)
    cell.font    = _font(bold=True, white=True)
    cell.border  = _border()
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + span - 1)
    return cell


def _data_cell(ws, row, col, value, bg=C_WHITE, bold=False, center=False) -> openpyxl.cell.Cell:
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill   = _fill(bg)
    cell.font   = _font(bold=bold)
    cell.border = _border()
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=True,
    )
    return cell


def _pct(v) -> str:
    return f"{v*100:.0f}%" if v is not None else "—"


def _delta_str(before, after) -> str:
    if before is None or after is None:
        return "—"
    diff = after - before
    sign = "+" if diff >= 0 else ""
    return f"{sign}{diff*100:.0f}pp"


def _delta_fill(before, after) -> str:
    if before is None or after is None:
        return C_WHITE
    if after > before:
        return C_PASS
    if after < before:
        return C_FAIL
    return C_NEUTRAL


# ── Data loaders ──────────────────────────────────────────────────────────────

def _load_results(run_name: str, system: str) -> list:
    path = HERE / f"results_{run_name}_{system}.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def _load_metrics(run_name: str, system: str) -> dict:
    path = HERE / f"metrics_{run_name}_{system}.json"
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _load_gt() -> list:
    return json.loads(GT_PATH.read_text(encoding="utf-8"))


# ── Sheet 1: Summary ──────────────────────────────────────────────────────────

def _sheet_summary(wb, systems: list, before_name: str, after_name: str):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False

    row = 1
    # Title
    cell = ws.cell(row=row, column=1, value="Narrative Validation Eval — Before / After Summary")
    cell.font = Font(bold=True, size=14, color=C_HEADER)
    ws.merge_cells(f"A{row}:L{row}")
    row += 2

    for system in systems:
        m_before = _load_metrics(before_name, system)
        m_after  = _load_metrics(after_name,  system) if after_name else {}

        ov_b = m_before.get("overall", {})
        ov_a = m_after.get("overall", {})

        # System header
        _header_cell(ws, row, 1, f"System: {system.upper()}", span=8, bg=C_SUBHEADER)
        row += 1

        # Column headers
        cols = ["Metric", "Before", "After", "Δ"] if after_name else ["Metric", "Value"]
        for ci, label in enumerate(cols, 1):
            _header_cell(ws, row, ci, label)
        row += 1

        metrics = [
            ("Criteria Recall",           ov_b.get("recall"),        ov_a.get("recall")),
            ("Criteria Precision",         ov_b.get("precision"),     ov_a.get("precision")),
            ("F1 Score",                   ov_b.get("f1"),            ov_a.get("f1")),
            ("Avg Compliance Score (0-10)",ov_b.get("avg_compliance_score"), ov_a.get("avg_compliance_score")),
            ("Avg Rewrite Faithfulness",   ov_b.get("avg_faithfulness"), ov_a.get("avg_faithfulness")),
        ]

        for label, vb, va in metrics:
            is_pct = vb is not None and isinstance(vb, float) and vb <= 1.0
            vb_str = _pct(vb) if is_pct else (str(vb) if vb is not None else "—")
            va_str = _pct(va) if is_pct else (str(va) if va is not None else "—")
            _data_cell(ws, row, 1, label, bold=True)
            _data_cell(ws, row, 2, vb_str, center=True)
            if after_name:
                _data_cell(ws, row, 3, va_str, center=True,
                           bg=_delta_fill(vb, va) if is_pct else C_WHITE)
                _data_cell(ws, row, 4, _delta_str(vb, va) if is_pct else "—",
                           center=True, bg=_delta_fill(vb, va) if is_pct else C_WHITE)
            row += 1

        row += 1  # spacer between systems

    ws.column_dimensions["A"].width = 30
    for col in "BCDE":
        ws.column_dimensions[col].width = 14


# ── Sheet 2: Per-Project ──────────────────────────────────────────────────────

def _sheet_per_project(wb, systems: list, before_name: str, after_name: str):
    ws = wb.create_sheet("Per-Project")
    ws.sheet_view.showGridLines = False

    gt = _load_gt()
    gt_map = {p["project_name"]: p for p in gt}

    row = 1
    cell = ws.cell(row=row, column=1, value="Per-Project Scores — Before / After")
    cell.font = Font(bold=True, size=13, color=C_HEADER)
    ws.merge_cells(f"A{row}:I{row}")
    row += 2

    for system in systems:
        res_b = {r["project_name"]: r for r in _load_results(before_name, system)}
        res_a = {r["project_name"]: r for r in _load_results(after_name,  system)} if after_name else {}

        _header_cell(ws, row, 1, f"System: {system.upper()}", span=9, bg=C_SUBHEADER)
        row += 1

        headers = ["Project", "Score Before", "Verdict Before"]
        if after_name:
            headers += ["Score After", "Verdict After", "Score Δ"]
        headers += ["Human Missing Criteria"]
        for ci, h in enumerate(headers, 1):
            _header_cell(ws, row, ci, h)
        row += 1

        for gt_proj in gt:
            name    = gt_proj["project_name"]
            missing = gt_proj["_summary"]["missing_criteria"]
            rb      = res_b.get(name, {})
            ra      = res_a.get(name, {}) if after_name else {}

            def _score(res):
                r = res.get("result") or {}
                return r.get("layer1", {}).get("compliance_score")

            def _verdict(res):
                r = res.get("result") or {}
                return r.get("overall_verdict", "—") if res.get("status") == "ok" else res.get("status", "—")

            sb = _score(rb); sa = _score(ra)
            vb = _verdict(rb); va = _verdict(ra)

            missing_labels = ", ".join(CRITERIA_LABELS[m] for m in missing) if missing else "All present"

            _data_cell(ws, row, 1, name, bold=True)
            _data_cell(ws, row, 2, sb if sb is not None else "—", center=True,
                       bg=VERDICT_FILL.get(vb, C_WHITE))
            _data_cell(ws, row, 3, vb, center=True, bg=VERDICT_FILL.get(vb, C_WHITE))
            col = 4
            if after_name:
                _data_cell(ws, row, col,   sa if sa is not None else "—", center=True,
                           bg=VERDICT_FILL.get(va, C_WHITE))
                _data_cell(ws, row, col+1, va, center=True, bg=VERDICT_FILL.get(va, C_WHITE))
                delta = (sa - sb) if (sa is not None and sb is not None) else None
                delta_str = (f"+{delta}" if delta and delta > 0 else str(delta)) if delta is not None else "—"
                delta_bg  = C_PASS if (delta and delta > 0) else (C_FAIL if (delta and delta < 0) else C_WHITE)
                _data_cell(ws, row, col+2, delta_str, center=True, bg=delta_bg)
                col = 7
            _data_cell(ws, row, col, missing_labels, bg=C_NEUTRAL if missing else C_PASS)
            row += 1

        row += 2

    ws.column_dimensions["A"].width = 38
    for i in range(2, 9):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions[get_column_letter(7 if after_name else 4)].width = 38


# ── Sheet 3: Per-Criterion ────────────────────────────────────────────────────

def _sheet_per_criterion(wb, systems: list, before_name: str, after_name: str):
    ws = wb.create_sheet("Per-Criterion")
    ws.sheet_view.showGridLines = False

    row = 1
    cell = ws.cell(row=row, column=1, value="Per-Criterion Metrics — Before / After")
    cell.font = Font(bold=True, size=13, color=C_HEADER)
    ws.merge_cells(f"A{row}:J{row}")
    row += 2

    for system in systems:
        m_before = _load_metrics(before_name, system)
        m_after  = _load_metrics(after_name,  system) if after_name else {}

        _header_cell(ws, row, 1, f"System: {system.upper()}", span=10, bg=C_SUBHEADER)
        row += 1

        headers = ["Criterion", "Recall Before", "Precision Before", "F1 Before"]
        if after_name:
            headers += ["Recall After", "Precision After", "F1 After",
                        "Δ Recall", "Δ Precision"]
        for ci, h in enumerate(headers, 1):
            _header_cell(ws, row, ci, h)
        row += 1

        pc_b = m_before.get("per_criterion", {})
        pc_a = m_after.get("per_criterion",  {})

        for c in CRITERIA_ORDER:
            cb = pc_b.get(c, {}); ca = pc_a.get(c, {})
            rb = cb.get("recall"); pb = cb.get("precision"); fb = cb.get("f1")
            ra = ca.get("recall"); pa = ca.get("precision"); fa = ca.get("f1")

            _data_cell(ws, row, 1, CRITERIA_LABELS[c], bold=True)
            _data_cell(ws, row, 2, _pct(rb), center=True, bg=C_FAIL if (rb is not None and rb < 0.5) else C_PASS if (rb is not None and rb >= 0.8) else C_WARN if rb is not None else C_WHITE)
            _data_cell(ws, row, 3, _pct(pb), center=True)
            _data_cell(ws, row, 4, _pct(fb), center=True)
            if after_name:
                _data_cell(ws, row, 5, _pct(ra), center=True, bg=_delta_fill(rb, ra))
                _data_cell(ws, row, 6, _pct(pa), center=True)
                _data_cell(ws, row, 7, _pct(fa), center=True)
                _data_cell(ws, row, 8, _delta_str(rb, ra), center=True, bg=_delta_fill(rb, ra))
                _data_cell(ws, row, 9, _delta_str(pb, pa), center=True, bg=_delta_fill(pb, pa))
            row += 1

        row += 2

    ws.column_dimensions["A"].width = 28
    for i in range(2, 11):
        ws.column_dimensions[get_column_letter(i)].width = 16


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Generate before/after Excel eval report.")
    parser.add_argument("--before",  default="baseline", help="Run name for 'before' results (default: baseline)")
    parser.add_argument("--after",   default=None,       help="Run name for 'after' results (omit for single-run report)")
    parser.add_argument("--system",  choices=["rag", "agent", "both"], default="both")
    args = parser.parse_args()

    systems = ["rag", "agent"] if args.system == "both" else [args.system]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default empty sheet

    _sheet_summary(wb,     systems, args.before, args.after)
    _sheet_per_project(wb, systems, args.before, args.after)
    _sheet_per_criterion(wb, systems, args.before, args.after)

    suffix = f"{args.before}_vs_{args.after}" if args.after else args.before
    out = HERE / f"eval_report_{suffix}.xlsx"
    wb.save(out)
    print(f"Wrote → {out}")


if __name__ == "__main__":
    main()